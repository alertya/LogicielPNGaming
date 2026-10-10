# -*- coding: utf-8 -*-
"""
Created on Fri Apr 14 12:05:00 2023

@author: USER
"""

import numpy as np
import random
import warnings
import time
import csv
import os
import sys
import math
import pandas as pd
import tkinter as tk
from tkinter import ttk, messagebox
import customtkinter as ctk
import utils
from config import *
import matplotlib.pyplot as plt
from tkinter import messagebox
from openpyxl import load_workbook
from ComputeFunction import Navigation, Maladie,FonctionEquipage,CalculMarchandise

from Display import *
#####Variable voyage

NbJour=6
CompetenceVigie=2

####Constantes
Nb=182
Nbsss=2
Superficie=106460000/2
NbDaysOfCyclone=59  ###https://public.wmo.int/fr/cyclones-tropicaux-0
NbDaysOfTempest=27
Surface=np.pi*750*750

SeuilMa=4
DeMa=20
SeuilAv=4
DeAv=40

Munitions=['Boulets','Rames','Mitraille']
Proba=round(Superficie/Surface*182/59)


def Choice_Rencontre(title: str, message: str, options: list[str]) -> str:
    """
    Create a CustomTkinter popup dialog with multiple buttons.

    :param title: Window title
    :param message: Text message displayed
    :param options: List of button labels
    :return: The label of the button clicked
    """
    # Demo app
    ctk.set_appearance_mode("System")
    ctk.set_default_color_theme("blue")

    app = ctk.CTk()
    app.title("CustomTkinter Multi-Button Dialog")
    app.geometry("600x400")
    choice = ctk.StringVar(value="")

    # Create popup
    popup = ctk.CTkToplevel()
    popup.title(title)
    popup.geometry("600x400")
    popup.grab_set()  # lock focus to popup

    # Message
    label = ctk.CTkLabel(popup, text=message, wraplength=300, justify="center")
    label.pack(pady=20, padx=10)

    # Button container
    btn_frame = ctk.CTkFrame(popup)
    btn_frame.pack(pady=10)

    # Create one button per option
    for opt in options:
        def handler(x=opt):  # capture opt in closure
            choice.set(x)
            popup.destroy()

        btn = ctk.CTkButton(btn_frame, text=opt, command=handler, width=80)
        btn.pack(side="left", padx=10)

    # Wait until popup closes
    popup.wait_window()
    return choice.get()


def CalculVoyage(NbJour,NbHommes):
    #####Variabe cout traversée
    Text=""
    CoutVivre = 1 / 150  ###Nourriture matlot par jour par personne
    CoutRhum = 10 / 150  ####Rhum par jour par matelot
    Solde = 0.31  ##Solde par jour par matelot
    CoutSabre = 1
    CoutMousquet = 10
    CoutMunition = 0.2  ##10 munitions par matelot
    NbMatelot = NbHommes
    ##########Variable tempete
    Nb = 182
    Nbsss = 2
    Superficie = 106460000 / 2
    NbDaysOfCyclone = 59  ###https://public.wmo.int/fr/cyclones-tropicaux-0
    NbDaysOfTempest = 27
    Surface = np.pi * 750 * 750
    Proba = round(Superficie / Surface * 182 / 59)
    CoutBase=CoutSabre*NbMatelot+CoutMousquet*NbMatelot+CoutMunition*NbMatelot
    CoutRation=CoutRhum*NbMatelot*NbJour+CoutVivre*NbMatelot*NbJour
    Text=Text+"\n Cout traversee : \n"+"Cout vivre :" + str(CoutVivre * NbJour * NbMatelot)+"\n Cout rhum :" + str(CoutRhum * NbJour * NbMatelot)+"\n CoutSabre : " + str(CoutSabre * NbMatelot)+"\nCoutFusil : " + str(CoutMousquet * NbMatelot)+"\nMunition : " + str(CoutMunition * NbMatelot)+"\nCout solde : " + str(Solde * NbJour * NbMatelot)+"\n \n"
    Text=Text+f" Cout ration {CoutRation} \n Cout base : {CoutBase} \n"
    return Text,CoutRation,CoutBase
    
def DistanceVision(Coeff,Nb):
    Dmax=(5+Coeff)/1.82
    Reus=utils.Test(Nb,0)
    if Reus<0:
        Max=utils.MinGauss(Dmax/2,Dmax/4,1)
    else:

        Max=utils.MaxGauss(Dmax/2,Dmax/4,Reus+1)
    Dist=Max
    Text="Situe a "+str(Dist)+" miles nautiques \n"
    Angle = round(random.uniform(0, 360), 2)
    Text =Text+ "A "+str(Angle)+" degrés du navire \n"
    return Text,Dist




def CheckNavire(Bateau):
    Garde_Cote=False
    DFTemp = pd.read_csv(BASE_PATH+"/RencontreNavire.csv", sep=";", decimal=",", encoding="cp1252")
    Lists = ['Brick', 'Brigantin', 'Corvette', 'Cotre', 'Deux-ponts trois-mâts barque', 'Deux-ponts trois-mâts carré',
             'Flibot', 'Flûte', 'Frégate trois-mâts barque', 'Frégate trois-mâts carré', 'Gabare',
             'Galiote (Brick PSx1.5)', 'Goélette balaou', 'Goélette brick', 'Goélette franche', 'Goélette de guerre',
             'Schooner', 'Houari (Goélette franche avec modif)', 'Ketch (Brigantin avec modif)', 'Langard (Brigantin)',
             'Négrier (Brick ou Marchand avec modif)', 'Paquebot (Goélette balaou)', 'pingre (Flûte)',
             'patach (Goélette franche)', 'pilote (Goélette franche, montre la route à suivre au port)',
             'pinnasse (Goelette franche)', 'plut (Brigantin, Néerlandais)',
             'prame (deux ponts trois mat barque avec modif)', 'ramberge (Sloop, riviere, anglais)', 'Schooner',
             'Senau (brick avec modif)', 'Sloop', 'Corvette', 'traversier (Chasse-marée, francais)',
             'trois-mâts goélette', 'trois-ponts', 'yacht (sloop)']
    if Bateau=="Garde côte (relancer jusqu’à trouver sloop, goélette, brigantin, brick, chebec, etc)":
        while Bateau not in Lists:
            de = random.randint(1, 100)
            DFTemp = pd.read_csv(BASE_PATH+"/RencontreNavire.csv", sep=";", decimal=",", encoding="cp1252")
            Bateau = DFTemp[DFTemp["DeBateau"] >= de]["Bateau"].iloc[0]
        Garde_Cote = True

    return Bateau,Garde_Cote



def Test(Nb):
    Success=0
    for k in range(Nb):
        Success=0
        if Nb==0:
            Test=random.choice(range(12))+1
            if Test>9:
                Success=-1
            elif Test==1:
                Success=2
            elif Test<=5:
                Success=1
        elif k>0:
            for h in range(k):
                Test=random.choice(range(10))+1
                if Test==1:
                    if Success==-1:
                        Success=0
                    else:
                        Success=Success+2
                elif Test<=5:
                    if Success==-1:
                        Success=0
                    Success=Success+1
                elif Test>9 and Success<1:
                    Success=-1
    return Success
def ReconnaissanceNavire(Vigilance,Navire):
    Te=0
    Text="Votre vigie a apercu un navire avec \n"
    Cat=['CategorieNavire','NbCanons','NbMats','Longueur']
    for k in Cat:
        Succes=utils.Test(Vigilance,0)
        Valeur = Navire[k]
        Estimation=utils.MaxGauss(Valeur,Valeur/2,Succes)
        Text+=" Un "+k+" de "+str(Estimation)+" \n"
    return Text

def RMarchand(SeuilMa,DeMa,SeuilAv,DeAv,Distance,CompVigie,Region,Name1,Zone,CompagnieCommerciale):
    ##########Variable tempete
    Nb = 182
    NbJour=1
    Actions=[""]
    Superficie = 106460000 / 2
    NbDaysOfCyclone = 59  ###https://public.wmo.int/fr/cyclones-tropicaux-0
    NbDaysOfTempest = 27
    Surface = np.pi * 750 * 750
    Proba = round(Superficie / Surface * Nb / NbDaysOfCyclone)
    NavireOrigine=pd.read_csv(os.path.join(NAVIRE_PATH, Name1), sep=";", decimal=",", encoding="cp1252")
    Text=f"Vous voyagez avec le navire {Name1} qui est un {NavireOrigine['Nom'].iloc[0]} \n"
    NbHommesNavire = NavireOrigine['Equipage'].iloc[0]
    Text=Text+"Vous avez un tonnage max de "+str(NavireOrigine['TonnageMax'].iloc[0])+ f" et {NbHommesNavire} hommes disponibles à bord \n"
    SalaireTotal=FonctionEquipage.AttributionSalaireParRole(NavireOrigine['EquipageNom'].iloc[0])
    Text=Text+"Le salaire journalier de votre equipage est de "+str(SalaireTotal/30)+ " \n"
    Text=Text+"La valeur de combat est de "+str(NavireOrigine['ValeurCombat'])+" \n"
    Tmp,CoutRation,CoutBase=CalculVoyage(NbJour,NbHommesNavire)
    Text = Text+Tmp+ "\n"
    CoutSalaire=0
    Interet=0
    while Distance>-1:
        CoutSalaire+=SalaireTotal/30
        CoutSalaire+=CoutRation
        Distance=Distance-NavireOrigine['Vitesse moyenne'].iloc[0]  ####On déduit la distance restante par la distance parcourue en moyenne par le navire
        DistanceErreur=Navigation.TestNavigation(Name1)

        Distance=Distance+DistanceErreur

        Text=Text+ " Rapport du jour de voyage numéro "+str(NbJour)+" \n \n \n"
        ####Test de maladie au sein de l'équipage
        Text=Text+Maladie.test_maladie(1,1,NavireOrigine['EquipageNom'].iloc[0])+" \n \n"
        DeAv=int(DeAv)
        ###Test de rencontre de récif par l'hydrographie du pilote
        Text =Text+Navigation.TestHydrographie(Name1, Zone)
        ###Test de rencontre de tempete
        Text = Text + Navigation.Tempest(Name1,Proba, NbJour)[0]+" \n \n"
        print("ANNONCE VOYAGE ")
        print(Text)
        print("FIN ANNONCE VOYAGE")
        if "coule" in Text.lower():
            Text = Text + "\n FIN DU VOYAGE"
            Distance = -999
        Distance = Distance + Navigation.Tempest(Name1, Proba, NbJour)[1]

        if random.choice(range(DeAv))+1<=SeuilAv:
            Hour=random.choice(range(12))+7
            Text = Text + "Rencontre d'un navire AVENTURIER effectué au bout de " + str(NbJour) + "jour a " + str(
                Hour) + " heure \n"
            Navire,Temp=RencontreNavire(str(NbJour)+"_AVENTURIER"+Region+"_"+Zone,Region,1,"Pirate")
            if "coule" in Temp or "COULE" in Temp:
                Text=Text+"\n FIN DU VOYAGE \n \n \n"
                Distance=-999
            Text = Text + Temp+"\n"
            Text1= ReconnaissanceNavire(CompVigie, Navire)
            Temp,Dist= DistanceVision(BonusVisionHauteurMat(NavireOrigine['CategorieNavire'].iloc[0]), CompVigie)
            Temp=Temp+Text1
            Interet = CalculInteretCombat(Navire, NavireOrigine)
            FormulaireRencontre(NavireOrigine, Navire, Dist, Temp,Interet)
            Text = Text + Temp + "\n"
            Text = Text + "FIN DE RENCONTRE DE NAVIRE \n \n"


            if Interet>0:
                print("Le navire cherche à vous piller")
                Text+=" \n LE NAVIRE CHERCHE A VOUS PILLER \n"
            messagebox.showinfo("Info",Temp)
            print("FIN DE RENCONTRE DE NAVIRE\n \n")
            print("")
            print("")


        if "N" in str(DeMa):
            DeMar = int(DeMa.replace('N', ''))
            for h in range(DeMar):
                Text=Text+"Rencontre d'un navire MARCHAND Numero "+str(h)+"effectué au bout de "+str(NbJour)+"jour "+ "\n"
                Navire,Temp=RencontreNavire(str(NbJour)+"_MARCHAND"+Region+"_"+Zone,Region,1,CompagnieCommerciale)
                Text=Text+Temp+"\n"
                Text1=ReconnaissanceNavire(CompVigie,Navire)
                Temp,Dist= DistanceVision(BonusVisionHauteurMat(NavireOrigine['CategorieNavire'].iloc[0]), CompVigie)
                Temp = Text1 + Temp
                FormulaireRencontre(NavireOrigine, Navire, Dist, Temp,Interet)
                Text = Text + Temp + "\n"
                messagebox.showinfo("Info", Temp)
                Text=Text+"FIN DE RENCONTRE DE NAVIRE \n \n"
                print("FIN DE RENCONTRE DE NAVIRE")
                print("")
                print("")
        else:
            DeMar=int(DeMa)
            if random.choice(range(DeMar))+1<=SeuilMa:

                Hour=random.choice(range(12))+7
                Text=Text+"Rencontre d'un navire MARCHAND effectué au bout de "+str(NbJour)+"jour a "+str(Hour)+ " heure \n"
                Navire,Temp=RencontreNavire(str(NbJour)+"_MARCHAND"+Region+"_"+Zone,Region,1,CompagnieCommerciale)
                Text=Text+Temp+"\n"
                Text1 =ReconnaissanceNavire(CompVigie,Navire)
                Temp,Dist= DistanceVision(BonusVisionHauteurMat(NavireOrigine['CategorieNavire'].iloc[0]), CompVigie)
                Temp=Text1+Temp
                Text = Text + Temp + "\n"
                messagebox.showinfo("Info", Temp)
                FormulaireRencontre(NavireOrigine, Navire, Dist, Temp,Interet)
                Text=Text+"FIN DE RENCONTRE DE NAVIRE \n \n"
                print("FIN DE RENCONTRE DE NAVIRE")
                print("")
                print("")

        NbJour += 1

    CoutTotal=CoutBase+CoutRation
    Text=Text+"\n Le voyage aura couté au total "+str(CoutSalaire)+"\n"
    Text=Text+"\n L'equipement aura couté au total "+str(CoutBase)+"\n"
    Text=Text+"\n Vous vous etes arrete à un port au niveau de la "+Region+"\n Vous pouvez consulter les données disponibles et leur court achat/revente dans e fichier : "+"CoursMarchandise"+Region+".csv"

    print('DUREE DU VOYAGE'+str(NbJour))
    #Marchandises=CalculMarchandise.CoursMarchandise(1,Region,"Au port (petit)")
    #Tmp,CAchat,CVente=CalculMarchandise.TrouverMarchand(2,Marchandises,Region)
    Text=Text+Tmp
    f = open(BASE_PATH+"/Rencontre/"+Region+"_"+Zone+".txt", "w")
    f.write(Text)
    f.close()
    return CoutBase,CoutSalaire,Text


#RMarchand(SeuilMa,DeMa,SeuilAv,DeAv,NbJour,CompetenceVigie)





def CalculCombatNavire(Navire):
    EsperanceParSucces=0.6*0.5
    ValCanon=float(Navire['Valeur canonnade'].iloc[0])
    Pointage=float(Navire['Pointage'].iloc[0])
    DegMoy=0.9925 * np.exp(0.2298 * (ValCanon - 9.9721 + 2 * Pointage*EsperanceParSucces))
    return Navire['StructureCoque'].iloc[0]*DegMoy/(7-float(Navire['Recharge'].iloc[0])*EsperanceParSucces)


def CalculInteretCombat(NavirePirate,NavirePJ):
    Interet=0
    SeuilCombatInteret=0.8   ###Ratio de valeur de combat pour lequel le pirate a un interet
    NavirePirate=NavirePirate.iloc[0]
    NavirePJ=NavirePJ.iloc[0]
    RatioCombat=NavirePirate['ValeurCombat']/ NavirePJ['ValeurCombat']
    SeuilMarchandise=30
    if RatioCombat>=SeuilCombatInteret:
        if (NavirePirate['TonnageMax']-NavirePirate['Tonnage'])>SeuilMarchandise/RatioCombat:
            Interet=1

    return Interet

def BonusVisionHauteurMat(CategorieNavire):
    if CategorieNavire==0: ##Si c'est une chaloupe, bonus de 5
        Bonus=5
    if CategorieNavire==1: ##Si c'est un sloop, bonus de 8
        Bonus=8
    if CategorieNavire==2: ##Si c'est une géoellte, bonus de 10
        Bonus=10
    if CategorieNavire==3: ##Si c'est un vaisseau de ligne, bonus de 18
        Bonus=18
    return Bonus


def ChoixRencontre(NavirePJ,NavirePNJ,Distance,Choix,Interet):
    CommerceCapitaine=utils.Compute(1)
    Text=""
    if Choix=="NeRienFaire":
        if Interet==0:
            Text+="Vous n'avez rien a faire, le navire s'éloigne de vous et vous continuez votre trajet"
        else:
            Text+="Le navire se rapproche de vous \n"

    LireEcrireCapitaine=utils.Compute(1)
    if Choix=="Poursuivre":
        Navigation.CoursePoursuite(3,Distance)
    ReussiteLire=utils.Test(LireEcrireCapitaine,0)
    if Choix=="Attaquer":
        Text+="Vous avez choisi d'attaquer le "+NavirePNJ['Name']+"\n Veuillez vous rendre dans l'onglet BatailleNavale et sélectionnez les navires combattants ainsi que vos munitions \n"
    ReussiteCommerce=utils.Test(CommerceCapitaine,0)
    if Choix=="Commercer":
        if "Pirate" in NavirePNJ['Allegence'] or "Interlope" in NavirePNJ['Allegence']:
            TauxVente=80-ReussiteCommerce*10
            TauxAchat=100+ReussiteCommerce*10
            Text+=("Vous pouvez acheter ou vendre des marchandises au capitaine, faites un text de Commerce sous Expression,"
                   " \n vous pouvez vendre à ")+str(TauxVente) +" plus votre reussite*10 % \n et acheter à "+str(TauxAchat)+"moins votre reussite*10 % \n"

        else:
            Result=random.randint(1,10)
            if Result<2:
                Text += "Le capitaine n'est pas intéressé pour commerce avec vous, vous continuez votre trajet séparemment \n"
            else:
                Text += "Le capitaine vous demande les papiers de provenance de vos marchandises avant de vous conseille le port le plus proche pour écouler votre matrchandise\n"
                if ReussiteLire>1:
                        Text+="\n si c'est un faux,  vous devez générer une bataille terrestre entre "+NavirePJ['EquipageNom']+ " et "+NavirePNJ['EquipageNom']+(""
                        " \n si vous perdez, la partie est terminée pour vous, si vous gagnez vous pouvez piller le navire et ou récupérer également le navire si il vous reste assez d'hommes \n")

    Text=Text+" \n si vous avez effectué, toutes les actions exigées ou souhaitées, vous pouvez fermer la fenêtre ! \n"
    messagebox.showinfo(
        "Résultat",
        Text
    )
    return Text


def DeterminationParametre(Periode,RegionMaritime):
    Regions = pd.read_csv(BASE_PATH + "/ListeRegions.csv", sep=";", encoding='cp1252', decimal=",")
    ZoneCompagnie=Regions[Regions["RegionsCommerciale"]==RegionMaritime]["RegionsCompagnie"].iloc[0]
    ZoneNavire = Regions[Regions["RegionsCommerciale"] == RegionMaritime]["RegionsRencontre"].iloc[0]
    Comp=random.randint(1,100)
    Parametres=pd.read_csv(BASE_PATH+'/CompagnieCommerciale.csv',sep=";",encoding='cp1252',decimal=",")
    Parametres=Parametres[Parametres['Zone']==ZoneCompagnie]
    Parametres = Parametres[Periode<=Parametres['PeriodMax'] ]
    Parametres=Parametres[Comp<=Parametres['Intervalle']]
    Parametres=Parametres.iloc[0]
    Parametres["Compagnie"]=Parametres["Acteur"]+Parametres["Nationalite"]
    Navire=random.randint(1,100)
    Navires=pd.read_csv(BASE_PATH+"/GeneNavire.csv",sep=";",encoding='cp1252',decimal=",")
    Navires = Navires[Navires['Zone'] == ZoneNavire]
    print(Navires)
    Navire=Navires[Navire<=Navires['DeBateau']]["Bateau"].iloc[0]
    Navire = CheckNavire(Navire)
    Parametres["Navire"]=Navire
    return Parametres


def Escale(RegionMaritime,Port,Navire,Action,SuccesAction):
    Navire=pd.read_csv(NAVIRE_PATH+Navire+".csv",sep=";",encoding="cp1252", decimal=",")
    Navire=Navire.iloc[0]
    ReparCoque=Navire['StructureCoqueMax']-Navire['StructureCoque']
    ReparVoile=Navire['StructureVoileMax']-Navire['StructureVoile']
    NbJourReparation=0
    NbJourEscale=0
    Text=""
    if Action=="Reparer":
        CompCharpentier=utils.Compute(1)
        Cout=(ReparCoque+ReparCoque)/(Navire['StructureCoqueMax']+Navire['StructureVoileMax'])*Navire['Cout']
        Navire['StructureCoque']=Navire['StructureCoqueMax']
        Navire['StructureVoile']=Navire['StructureVoileMax']
        for k in range(ReparCoque+ReparVoile):
            NbJourReparation+=utils.TestValeurNonNumerique(CompCharpentier,0)[0]
        NbJourEscale=max(NbJourReparation,NbJourEscale)
        Navire.to_csv(NAVIRE_PATH+Navire+".csv",sep=";",encoding="cp1252", decimal=",")
        Text+="Votre navire est réparé pour un cout de "+str(Cout)+" pièces de huit en "+str(NbJourEscale)+ "jours \n"
    if Action=="Recruter":
        NbHommes=0
        Navire=RencontreNavire(1,RegionMaritime,0,"Defaut")[0].iloc[0]
        NbHommes+=Navire['Equipage']
        NbHommesRecrute=NbHommes/10*SuccesAction
        Text+=FonctionEquipage.Recrute(NbHommesRecrute,"Matelot",Navire,0)


    return Text



def FicheNavire(Name,Navire,Marchandises):
    # Ouvrir le fichier
    wb = load_workbook(BASE_PATH+ "/FicheNavire.xlsx")

    # Sélectionner la feuille
    ws = wb["Fiche"]
    Navire=Navire.iloc[0]
    # Remplir des cellules
    ws["B2"] = Name
    ws["D3"] = Navire['RegionsDepart']
    ws["D2"] = Navire['Allegence']
    ws["B4"] = Navire['Nom']
    ws["D4"] = Navire['Pirate']
    ws["G2"] = Navire['Manoeuvre']
    ws["G4"] = Navire['Combat']
    ws["G5"] = Navire['Pointage']
    ws["G6"] = Navire['Recharge']
    ws["G7"] = Navire['Ruse']
    ws["G8"] = Navire['ValeurCombat']
    ws["B13"] = Navire['NbCanons']
    ws["B12"] = Navire['Valeur canonnade']
    ws["B10"] = Navire['StructureCoque']
    ws["B11"] = Navire['StructureVoile']
    ws["D5"] = Navire['EquipageNom']
    ws["G10"] = Navire['Longueur']
    ws["G11"] = Navire['NbMats']
    ws["G12"] = Navire['CategorieNavire']
    ws["G13"] = Navire['Longueur']
    ws["B16"] = Navire['Pres']
    ws["B17"] = Navire['Largue']
    ws["B18"] = Navire['Grand largue']
    ws["B19"] = Navire['Vent arriere']
    ws["B20"] = Navire['Vitesse moyenne']
    ws["B21"] = Navire['TonnageMax']
    ws["B22"] = Navire['Tonnage']
    ws["B7"] = Navire['Equipage']
    ws["B6"] = Navire['Reputation']
    row=26
    for key, value in Marchandises.items():
        ws.cell(row=row, column=1).value = key
        ws.cell(row=row, column=2).value = value
        row+=1
    # Sauvegarder
    wb.save(BASE_PATH+"/Rencontre/"+Name+".xlsx")


def FormulaireRencontre(NavirePJ, NavirePNJ, Distance, Text,Interet):

    root1 = tk.Toplevel()
    root1.title("FormulaireRencontre")
    root1.geometry("600x200")

    # Texte en haut
    label_text = tk.Label(
        root1,
        text=Text,
        wraplength=550,
        justify="left"
    )
    label_text.pack(pady=20)

    # Cadre contenant les boutons
    frame_actions = tk.Frame(root1)
    frame_actions.pack(pady=10)

    Actions = [
        "NeRienFaire",
        "Poursuivre",
        "Combattre",
        "Commercer"
    ]

    def executer_choix(choix):

        TextResultat = ChoixRencontre(
            NavirePJ,
            NavirePNJ,
            Distance,
            choix,Interet
        )

        messagebox.showinfo(
            "Résultat",
            TextResultat
        )

        root1.destroy()

    # Création des boutons sur une seule ligne
    for k in Actions:

        bouton = tk.Button(
            frame_actions,
            text=k,
            command=lambda choix=k: executer_choix(choix)
        )

        bouton.pack(
            side="left",
            padx=5
        )

    # Bloque les interactions avec les autres fenêtres
    root1.grab_set()

    # Attend que la fenêtre soit fermée
    root1.wait_window()








































